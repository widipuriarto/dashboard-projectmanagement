import pandas as pd
import numpy as np
import openpyxl
import os
import re

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_PATH = os.path.join(BASE_DIR, 'database', 'Full Copy of Copy of PROJECT MANAGEMENT RSO 1 REG 3.xlsx')
OUTPUT_PATH = os.path.join(BASE_DIR, 'seed_data.csv')
DB_OUTPUT_PATH = os.path.join(BASE_DIR, 'database', 'seed_data.csv')
SQL_MIGRATION_PATH = os.path.join(BASE_DIR, 'database', 'migrate_dates.sql')

MONTH_MAP = {
    'jan': 1, 'januari': 1, 'january': 1,
    'feb': 2, 'februari': 2, 'february': 2, 'febuari': 2,
    'mar': 3, 'maret': 3, 'march': 3,
    'apr': 4, 'april': 4,
    'mei': 5, 'may': 5,
    'jun': 6, 'juni': 6, 'june': 6,
    'jul': 7, 'juli': 7, 'july': 7,
    'agu': 8, 'ags': 8, 'aug': 8, 'agustus': 8, 'august': 8,
    'sep': 9, 'sept': 9, 'september': 9,
    'okt': 10, 'oct': 10, 'oktober': 10, 'october': 10,
    'nov': 11, 'nop': 11, 'november': 11,
    'des': 12, 'dec': 12, 'desember': 12, 'december': 12
}

MONTH_REGEX = '|'.join(sorted(MONTH_MAP.keys(), key=len, reverse=True))

def parse_month_year_to_date(val):
    if pd.isna(val) or not str(val).strip():
        return None
    val_str = str(val).strip()
    
    # 1. Direct standard datetime parsing
    try:
        dt = pd.to_datetime(val_str, errors='coerce')
        if pd.notna(dt):
            return dt.strftime('%Y-%m-01')
    except Exception:
        pass
        
    # 2. Regex: Month + Year (e.g. 'Juli 2025', 'Sept 2023', 'Apr 2026', 'Mei 2025 (CO April)')
    match = re.search(rf'({MONTH_REGEX})\b.*?(\d{{4}})', val_str, flags=re.IGNORECASE)
    if match:
        m_str, y_str = match.group(1).lower(), match.group(2)
        month_num = MONTH_MAP.get(m_str)
        if month_num:
            return f'{y_str}-{month_num:02d}-01'
            
    # 3. Regex: Year + Month
    match_rev = re.search(rf'(\d{{4}}).*?({MONTH_REGEX})\b', val_str, flags=re.IGNORECASE)
    if match_rev:
        y_str, m_str = match_rev.group(1), match_rev.group(2).lower()
        month_num = MONTH_MAP.get(m_str)
        if month_num:
            return f'{y_str}-{month_num:02d}-01'
            
    return None

def generate_sql_migration(df, sql_path):
    bulan_cases = []
    komitmen_cases = []
    
    for _, row in df.iterrows():
        p_id = row['project_id_pm']
        if not p_id or pd.isna(p_id):
            continue
        p_id_esc = str(p_id).replace("'", "''")
        
        b_val = row['bulan_bc']
        if pd.notna(b_val) and str(b_val).strip():
            bulan_cases.append(f"    WHEN '{p_id_esc}' THEN '{b_val}'::date")
            
        k_val = row['komitmen_bc']
        if pd.notna(k_val) and str(k_val).strip():
            komitmen_cases.append(f"    WHEN '{p_id_esc}' THEN '{k_val}'::date")

    sql_content = f"""-- Auto-generated Database Migration Script by etl_script.py
-- Description: Normalize bulan_bc & komitmen_bc to standard DATE and alter column schema

BEGIN;

-- 1. Update bulan_bc and komitmen_bc across existing projects
UPDATE public.projects
SET 
  bulan_bc = CASE project_id_pm
{chr(10).join(bulan_cases)}
    ELSE NULL
  END,
  komitmen_bc = CASE project_id_pm
{chr(10).join(komitmen_cases)}
    ELSE NULL
  END;

-- 2. Alter column types to DATE
ALTER TABLE public.projects 
  ALTER COLUMN bulan_bc TYPE DATE USING bulan_bc::date,
  ALTER COLUMN komitmen_bc TYPE DATE USING komitmen_bc::date;

-- 3. Ensure proper table privileges for Supabase roles
GRANT ALL ON TABLE public.projects TO anon, authenticated, service_role;

COMMIT;
"""
    with open(sql_path, 'w', encoding='utf-8') as f:
        f.write(sql_content)
    print(f"[SQL MIGRATION] File migrasi SQL berhasil dibuat: {sql_path}")

def run_etl():
    print("\n=== Memulai proses ETL (Extract, Transform, Load) ===")
    
    try:
        # 1. AUTO-DETECT HEADER ROW
        df_test = pd.read_excel(FILE_PATH, sheet_name='new POPULASI PROJECT', header=0)
        header_idx = 0
        
        # Cek apakah 'ID PROJECT' ada di baris pertama
        if not any('ID PROJECT' in str(c).upper() or 'NAMA PROYEK' in str(c).upper() for c in df_test.columns):
            print("[INFO] Header tidak di baris 1, mencoba baris 2...")
            df_test = pd.read_excel(FILE_PATH, sheet_name='new POPULASI PROJECT', header=1)
            header_idx = 1
            
        df = df_test.copy()
        # Normalisasi ke UPPERCASE untuk menghindari error kapitalisasi
        df.columns = df.columns.astype(str).str.strip().str.upper()
        
        print(f"\n[DIAGNOSTIK] Kolom yang berhasil dibaca dari Excel (Header Row {header_idx + 1}):")
        print(df.columns.tolist())
        print("\n")
        
        print("2. Mengekstrak target URL asli (Hyperlink) menggunakan Openpyxl...")
        wb = openpyxl.load_workbook(FILE_PATH, data_only=True)
        ws = wb['new POPULASI PROJECT']
        
        # Ekstrak LINK BAST dari sheet APROVAL PJM NCX (Restorasi Logika)
        bast_mapping = {}
        if 'APROVAL PJM NCX' in wb.sheetnames:
            print("   -> Membangun relasi hyperlink BAST dari sheet 'APROVAL PJM NCX'...")
            ws_approval = wb['APROVAL PJM NCX']
            
            # Mapping header dari APROVAL PJM NCX
            approval_headers = {str(cell.value).strip().upper(): idx + 1 for idx, cell in enumerate(ws_approval[1]) if cell.value}
            idx_id_pm_bast = approval_headers.get('ID PM')
            idx_link_dok_bast = approval_headers.get('LINK DOK (TEAMS)')
            
            if idx_id_pm_bast and idx_link_dok_bast:
                for row in range(2, ws_approval.max_row + 1):
                    cell_id = ws_approval.cell(row=row, column=idx_id_pm_bast)
                    cell_link = ws_approval.cell(row=row, column=idx_link_dok_bast)
                    
                    if cell_id.value is not None and cell_link.hyperlink is not None:
                        bast_mapping[str(cell_id.value).strip()] = cell_link.hyperlink.target
        else:
            print("   -> [WARNING] Sheet 'APROVAL PJM NCX' tidak ditemukan untuk LINK BAST.")

        # Ekstrak LINK BAPL dari file eksternal Copy 2 BAPL.xlsx
        bapl_mapping = {}
        try:
            bapl_file_path = os.path.join(BASE_DIR, 'database', 'Copy 2 BAPL.xlsx')
            if os.path.exists(bapl_file_path):
                print("   -> Membangun relasi hyperlink BAPL dari 'Copy 2 BAPL.xlsx'...")
                wb_bapl = openpyxl.load_workbook(bapl_file_path, data_only=True)
                if 'BAPL' in wb_bapl.sheetnames:
                    ws_bapl = wb_bapl['BAPL']
                    
                    # Mencari index kolom (asumsi baris 1 adalah header)
                    bapl_headers = {str(cell.value).strip().upper(): idx + 1 for idx, cell in enumerate(ws_bapl[1]) if cell.value}
                    idx_id_pm_bapl = bapl_headers.get('ID PM', 2) # Kolom B
                    idx_link_teams = bapl_headers.get('LINK TEAMS', 16) # Kolom P
                    
                    # Cukup baca sampai baris 312 sesuai instruksi untuk optimasi
                    max_row_bapl = min(312, ws_bapl.max_row)
                    
                    for row in range(2, max_row_bapl + 1):
                        cell_id = ws_bapl.cell(row=row, column=idx_id_pm_bapl)
                        cell_link = ws_bapl.cell(row=row, column=idx_link_teams)
                        
                        if cell_id.value is not None and cell_link.hyperlink is not None:
                            bapl_mapping[str(cell_id.value).strip()] = cell_link.hyperlink.target
        except Exception as e:
            print(f"   -> [WARNING] Gagal membaca Copy 2 BAPL.xlsx: {e}")

        # Penyesuaian baris Openpyxl
        ws_header_row = header_idx + 1
        
        # Mapping header openpyxl ke UPPERCASE
        headers = {str(cell.value).strip().upper(): idx + 1 for idx, cell in enumerate(ws[ws_header_row]) if cell.value}
        idx_id_project_pm = headers.get('ID PROJECT PM')
        idx_bast = headers.get('LINK BAST')
        idx_bapl = headers.get('LINK BAPL')
        
        if 'LINK BAST' not in df.columns: df['LINK BAST'] = np.nan
        if 'LINK BAPL' not in df.columns: df['LINK BAPL'] = np.nan
        
        for i in range(len(df)):
            excel_data_row = i + ws_header_row + 1
            
            # Memasukkan LINK BAST dari mapping eksternal atau fallback
            link_bast_found = False
            if idx_id_project_pm is not None:
                cell_id_project = ws.cell(row=excel_data_row, column=idx_id_project_pm)
                if cell_id_project.value is not None:
                    id_val = str(cell_id_project.value).strip()
                    if id_val in bast_mapping:
                        df.at[i, 'LINK BAST'] = bast_mapping[id_val]
                        link_bast_found = True
                        
            # Jika tidak ada di mapping, gunakan URL dari file utama (jika ada)
            if not link_bast_found and idx_bast is not None:
                cell_bast = ws.cell(row=excel_data_row, column=idx_bast)
                if cell_bast.hyperlink:
                    df.at[i, 'LINK BAST'] = cell_bast.hyperlink.target
                    
            # Memasukkan LINK BAPL dari mapping eksternal atau fallback
            link_bapl_found = False
            if idx_id_project_pm is not None:
                cell_id_project = ws.cell(row=excel_data_row, column=idx_id_project_pm)
                if cell_id_project.value is not None:
                    id_val = str(cell_id_project.value).strip()
                    if id_val in bapl_mapping:
                        df.at[i, 'LINK BAPL'] = bapl_mapping[id_val]
                        link_bapl_found = True
                        
            # Jika tidak ada di mapping, gunakan URL dari file utama (jika ada)
            if not link_bapl_found and idx_bapl is not None:
                cell_bapl = ws.cell(row=excel_data_row, column=idx_bapl)
                if cell_bapl.hyperlink:
                    df.at[i, 'LINK BAPL'] = cell_bapl.hyperlink.target
                    
    except Exception as e:
        print(f"Error memproses file Excel: {e}")
        return

    print("3. Menerapkan Logic Fallback & Transformasi (Anti-KeyError)...")
    
    new_seg = df.get('NEW SEGMENT')
    old_seg = df.get('SEGMENT')
    if new_seg is not None and old_seg is not None:
        df['SEGMENT_FINAL'] = new_seg.fillna(old_seg)
    elif new_seg is not None:
        df['SEGMENT_FINAL'] = new_seg
    elif old_seg is not None:
        df['SEGMENT_FINAL'] = old_seg
    else:
        df['SEGMENT_FINAL'] = 'Uncategorized'
        
    new_witel = df.get('WITEL NEW')
    old_witel = df.get('WITEL')
    if new_witel is not None and old_witel is not None:
        df['WITEL_FINAL'] = new_witel.fillna(old_witel)
    elif new_witel is not None:
        df['WITEL_FINAL'] = new_witel
    elif old_witel is not None:
        df['WITEL_FINAL'] = old_witel
    else:
        df['WITEL_FINAL'] = 'Uncategorized'
    
    # MAPPING MENGGUNAKAN UPPERCASE
    column_mapping = {
        'NOMOR': 'nomor',
        'ID PROJECT PM': 'project_id_pm',
        'SEGMENT_FINAL': 'segment',
        'WITEL_FINAL': 'witel',
        'NAMA PROYEK': 'nama_proyek',
        'NAMA MITRA': 'nama_mitra',
        'NAMA PELANGGAN': 'nama_pelanggan',
        'IDLOP': 'idlop',
        'JENIS KONTRAK': 'jenis_kontrak',
        'NOMOR KL/SP/WO/NOPES': 'nomor_kl_sp_wo_nopes',
        'TGL KONTRAK': 'tgl_kontrak',
        'TAHUN KONTRAK': 'tahun_kontrak',
        'NILAI KONTRAK (OBL)': 'nilai_kontrak_obl',
        'QUOTE NUMBER': 'quote_number',
        'STATUS EMAIL HANDOVER': 'status_email_handover',
        'NOTE EMAIL HAND OVER': 'note_email_handover',
        'STATUS DOKUMEN OBL': 'status_dokumen_obl',
        'NO ORDER': 'no_order',
        'SID': 'sid',
        'JENIS LAYANAN': 'jenis_layanan',
        'STATUS DELIVERY PROJECT': 'status_delivery_project',
        'JENIS DOK DELIVERY': 'jenis_dok_delivery',
        'STATUS DOKUMEN DELIVERY': 'status_dokumen_delivery',
        'ID DIGIREVIEW': 'id_digireview',
        'TANGGAL REQUEST DR': 'tanggal_request_dr',
        'TANGGAL CLOSE DR': 'tanggal_close_dr',
        'NO BAST TELKOM': 'no_bast_telkom',
        'TANGGAL BAST': 'tanggal_bast',
        'TAHUN BAST': 'tahun_bast',
        'PM': 'pm',
        'WAKTU PENYELESAIAN BAST (HARI)': 'waktu_penyelesaian_bast',
        'KET UPDATE': 'ket_update',
        'STATUS ORDER': 'status_order',
        'BULAN BC': 'bulan_bc',
        'KOMITMEN BC': 'komitmen_bc',
        'LINK BAST': 'file_bast',
        'LINK BAPL': 'file_bapl',
        'STATUS GR': 'status_gr',
        'HARDFILE BAST': 'hardfile_bast'
    }
    
    for excel_col in column_mapping.keys():
        if excel_col not in df.columns:
            print(f" [WARNING] Kolom '{excel_col}' tidak ditemukan di Excel. Diisi kosong.")
    
    df = df.rename(columns=column_mapping)
    db_columns = list(column_mapping.values())
    
    for col in db_columns:
        if col not in df.columns:
            df[col] = np.nan
            
    df_final = df[db_columns].copy()
    
    df_final = df_final.dropna(subset=['project_id_pm', 'nama_proyek'], how='all')
    
    date_columns = ['tgl_kontrak', 'tanggal_request_dr', 'tanggal_close_dr', 'tanggal_bast']
    for date_col in date_columns:
        df_final[date_col] = pd.to_datetime(df_final[date_col], errors='coerce').dt.strftime('%Y-%m-%d')
        
    # Sanitasi & standardisasi kolom bulan_bc dan komitmen_bc ke format DATE 'YYYY-MM-01'
    df_final['bulan_bc'] = df_final['bulan_bc'].apply(parse_month_year_to_date)
    df_final['komitmen_bc'] = df_final['komitmen_bc'].apply(parse_month_year_to_date)

    df_final['status_email_handover'] = df_final['status_email_handover'].notna() & (df_final['status_email_handover'] != '')

    # FORMAT TAHUN (Hapus .0 akibat float NaN)
    for year_col in ['tahun_kontrak', 'tahun_bast']:
        df_final[year_col] = pd.to_numeric(df_final[year_col], errors='coerce').astype('Int64')

    # SANITASI DATA NUMERIK (Mencegah error 'invalid input syntax for type numeric/integer' di Supabase)
    # Beberapa cell memiliki teks ganda seperti "Tahap 1 : 536... Tahap 2: ..." atau "#VALUE!".
    for num_col in ['nomor', 'nilai_kontrak_obl', 'waktu_penyelesaian_bast']:
        df_final[num_col] = pd.to_numeric(df_final[num_col], errors='coerce').astype('Int64')

    # PEMBERSIHAN MULTILINE: Ganti semua karakter Enter (\n, \r) di dalam teks sel menjadi spasi biasa.
    df_final = df_final.replace(r'\r+|\n+', ' ', regex=True)
    
    # Simpan CSV ke direktori root dan database
    df_final.to_csv(OUTPUT_PATH, index=False, na_rep='')
    os.makedirs(os.path.dirname(DB_OUTPUT_PATH), exist_ok=True)
    df_final.to_csv(DB_OUTPUT_PATH, index=False, na_rep='')
    
    # Generate SQL Migration script untuk eksekusi langsung di Supabase SQL Editor
    generate_sql_migration(df_final, SQL_MIGRATION_PATH)

    print(f"\n[SUKSES] Seluruh data dan URL asli telah diamankan.")
    print(f"Data CSV siap di-import: {DB_OUTPUT_PATH} ({len(df_final)} baris)")
    print(f"Script migrasi SQL: {SQL_MIGRATION_PATH}")

if __name__ == '__main__':
    run_etl()

